"""
Downloads National Health Mission district Records of Proceedings
(RoPs) and extracts each district's sanctioned National Tobacco Control
Programme (NTCP) budget.

Actual NTCP expenditure is not published below state level; the RoPs
are the finest public figure, and they are approvals ("budget
sanctioned"), not money spent. Every RoP lists NTCP under NCD.4 with
the same FMR activity codes:

    104  Implementation of COTPA 2003 (enforcement squads, IEC, training)
    105  Implementation of ToFEI guidelines (tobacco-free schools)
    106  Tobacco cessation (cessation centres, medicines, training)

Only these activity lines are summed. NTCP staff salaries sit in each
RoP's human-resources section under a different code, and are left out
so that states are compared like for like.

States covered, and why the others are not (yet):

    Haryana       FY 2024-25, 22 district PDFs (layout text)
    Sikkim        FY 2024-25, one workbook; 6 current districts are
                  summed into NFHS-5's 4 (Pakyong was carved from East,
                  Soreng from West, in 2021)
    Meghalaya     FY 2024-25, one workbook with a state (SPMU) column
    Uttarakhand   FY 2024-25, 13 district PDFs
    Odisha        FY 2024-25 (from 2024-26 approvals), 30 district PDFs
    Rajasthan     FY 2024-25 (from 2024-26 approvals), 50 district PDFs
                  (the 2023 districts) folded into NFHS-5's 33; Dholpur's
                  file was published without its amount columns
    Assam         FY 2024-25, one NTCP budget sheet printed as a single
                  wide table across 14 pages, with a state HQ column

    Andhra Pradesh is left out: its only district RoPs are FY 2026-27
    scans (OCR text, figures that miss their own totals), and most of
    its 26 post-2022 districts straddle two of NFHS-5's 13.

District mapping. RoPs follow current districts, NFHS-5 the 2019-21
ones. A newer district is added to the NFHS-5 district it was carved
from. Where it came from two or more (six Rajasthan districts), it goes
wholly to the parent that held most of its Census 2011 population, and
every NFHS-5 district involved carries a mapping_note saying so.

Outputs (data/reference/oral_cancer/):
    ntcp_district_budget.csv  one row per NFHS-5 district
    ntcp_state_split.csv      state-HQ vs district share, where the
                              RoP shows both (Sikkim, Meghalaya, Assam)
"""

import csv
import io
import re
import time
from pathlib import Path
from urllib.parse import urljoin

import openpyxl
import pdfplumber
import requests
from pypdf import PdfReader

OUT_DIR = (
    Path(__file__).resolve().parent.parent
    / "data" / "reference" / "oral_cancer"
)

DISTRICT_PATH = OUT_DIR / "ntcp_district_budget.csv"
SPLIT_PATH = OUT_DIR / "ntcp_state_split.csv"

HEADERS = {"User-Agent": "Mozilla/5.0"}

CODES = ("104", "105", "106")

# --- Haryana -----------------------------------------------------------

HARYANA_PAGE = "https://nhmharyana.gov.in/page?id=284"

# RoP file name -> NFHS-5 district name, where they differ.
HARYANA_NAMES = {
    "CharkhiDadri": "Charkhi Dadri",
    "Gurugram": "Gurgaon",
    "Narnaul": "Mahendragarh",
}

# --- Sikkim ------------------------------------------------------------

SIKKIM_PAGE = "https://nhmsikkim.org/pip-rops/rops"

# Current district -> NFHS-5 district it was part of in 2019-21.
SIKKIM_NAMES = {
    "Gangtok": "East District",
    "Pakyong": "East District",
    "Gyalshing": "West District",
    "Soreng": "West District",
    "Mangan": "North  District",
    "Namchi": "South District",
}

# --- Meghalaya ---------------------------------------------------------

MEGHALAYA_PAGE = "https://nhmmeghalaya.nic.in/ROPs.html"

MEGHALAYA_NAMES = {
    "EKH": "East Khasi Hills",
    "WKH": "West Khasi Hills",
    "SWKH": "South West Khasi Hills",
    "WJH": "West Jaintia Hills",
    "EJH": "East Jantia Hills",
    "RB": "Ribhoi",
    "WGH": "West Garo Hills",
    "SWGH": "South West Garo Hills",
    "EGH": "East Garo Hills",
    "NGH": "North Garo Hills",
    "SGH": "South Garo Hills",
}

# --- Odisha ------------------------------------------------------------

ODISHA_PAGE = "https://nhmodisha.gov.in/pips/"

ODISHA_AJAX = "https://nhmodisha.gov.in/wp-admin/admin-ajax.php"

# Approved district plans for 2024-26 (Odisha publishes its district
# RoPs as "PIP"s headed "Approval 2024-26"). File stem -> NFHS-5 name;
# the listing also carries hospitals, medical colleges and city health
# societies, which are left out.
ODISHA_NAMES = {
    "Angul": "Anugul", "Balasore": "Baleshwar", "Baragarh": "Bargarh",
    "Bhadrak": "Bhadrak", "Bolangir": "Balangir", "Boudh": "Baudh",
    "Cuttack": "Cuttack", "Deogarh": "Debagarh", "Dhenkanal": "Dhenkanal",
    "Gajapati": "Gajapati", "Ganjam": "Ganjam",
    "Jagatsingpur": "Jagatsinghapur", "Jajapur": "Jajapur",
    "Jharsuguda": "Jharsuguda", "Kalahandi": "Kalahandi",
    "Kandhamal": "Kandhamal", "Kendrapada": "Kendrapara",
    "Keonjhar": "Kendujhar", "Khurda": "Khordha", "Koraput": "Koraput",
    "Malkangiri": "Malkangiri", "Mayurbhanj": "Mayurbhanj",
    "Nawrangpur": "Nabarangapur", "Nayagarh": "Nayagarh",
    "Nuapada": "Nuapada", "Puri": "Puri", "Rayagada": "Rayagada",
    "Sambalpur": "Sambalpur", "Sonepur": "Subarnapur",
    "Sundargarh": "Sundargarh",
}

# --- Uttarakhand -------------------------------------------------------

UTTARAKHAND_PAGE = "https://nhm.uk.gov.in/document-category/nhm-rop/"

# District name printed on the RoP cover -> NFHS-5 district name.
UTTARAKHAND_NAMES = {
    "Haridwar": "Hardwar",
    "Pauri Garhwal": "Garhwal",
}

# --- Rajasthan ---------------------------------------------------------

RAJASTHAN_PAGE = "https://rajswasthya.rajasthan.gov.in/link/districtrop.php"

# RoP file name (before " NHM"/" DROP"/" ROP") -> NFHS-5 district. The
# 2023 districts go to the NFHS-5 district they were carved from.
RAJASTHAN_NAMES = {
    "Ajmer": "Ajmer", "Alwar": "Alwar", "Anupgarh": "Ganganagar",
    "Balotra": "Barmer", "Banswara": "Banswara", "BARAN": "Baran",
    "Barmer": "Barmer", "Beawar": "Ajmer", "Bharatpur": "Bharatpur",
    "Bhilwara": "Bhilwara", "Bikaner": "Bikaner", "Bundi": "Bundi",
    "Chittorgarh": "Chittaurgarh", "Churu": "Churu", "Dausa": "Dausa",
    "Deedwana Kuchaman": "Nagaur", "Deeg": "Bharatpur",
    "Dholpur": "Dhaulpur", "Dudu": "Jaipur", "Dungarpur": "Dungarpur",
    "Gangapur City": "Sawai Madhopur", "Hanumangarh": "Hanumangarh",
    "Jaipsalmer": "Jaisalmer", "Jaipur 1": "Jaipur", "Jaipur 2": "Jaipur",
    "Jalore": "Jalor", "Jhalawar": "Jhalawar", "Jhunjhunu": "Jhunjhunun",
    "Jodhpur Rural": "Jodhpur", "Jodhpur Urban": "Jodhpur",
    "Karauli": "Karauli", "Kekri": "Ajmer", "Kherthal Tijara": "Alwar",
    "Kota": "Kota", "Kotputli Beharor": "Alwar", "Nagaur": "Nagaur",
    "Neem Ka Thana": "Sikar", "Pali": "Pali", "Phalodi": "Jodhpur",
    "Pratapgarh": "Pratapgarh", "Rajsamand": "Rajsamand",
    "Salumbar": "Udaipur", "Sanchore": "Jalor",
    "Sawai Madhopur": "Sawai Madhopur", "Shahpura": "Bhilwara",
    "Sikar": "Sikar", "Sirohi": "Sirohi", "Sri Ganganagar": "Ganganagar",
    "Tonk": "Tonk", "Udaipur": "Udaipur",
}

# 2023 districts carved from more than one NFHS-5 district: share of
# their Census 2011 population from each (tehsil populations, grouped
# by 2011 sub-district code; citypopulation.de/en/india/rajasthan/admin).
# Each goes wholly to its largest parent above.
RAJASTHAN_SPLITS = {
    "Anupgarh": {"Ganganagar": 0.80, "Bikaner": 0.20},
    "Beawar": {"Ajmer": 0.52, "Pali": 0.39, "Bhilwara": 0.09},
    "Gangapur City": {"Sawai Madhopur": 0.56, "Karauli": 0.44},
    "Kekri": {"Ajmer": 0.77, "Tonk": 0.23},
    "Kotputli Beharor": {"Alwar": 0.54, "Jaipur": 0.46},
    "Neem Ka Thana": {"Sikar": 0.63, "Jhunjhunun": 0.37},
}

# Published without the amount columns (labels only).
RAJASTHAN_UNREADABLE = {"Dholpur"}

# --- Assam -------------------------------------------------------------

ASSAM_SHEET = (
    "https://nhm.assam.gov.in/sites/default/files/swf_utility_folder/"
    "departments/nhm_lipl_in_oid_6/menu/document/"
    "ntcp_nhm_district_rop_budget_sheet_2024-26.pdf"
)

# Sheet heading -> NFHS-5 district. Bajali was carved from Barpeta.
ASSAM_NAMES = {
    "Bajali": "Barpeta",
    "Bongaingaon": "Bongaigaon",
    "Charaido": "Charaideo",
    "Dima Hassao": "Dima Hasao",
    "Kamrup (Metro)": "Kamrup Metropolitan",
    "Sivsagar": "Sivasagar",
    "South Salmara": "South Salmara Mancachar",
}

# Assam numbers every non-COTPA line 106.x; its own programme labels put
# 106.1-106.3 (school orientation, verification, college campaigns)
# under ToFEI, i.e. the national FMR 105.
ASSAM_TOFEI = {"106.1", "106.2", "106.3"}


def get(url, timeout=300, attempts=4):

    # State NHM servers (Haryana especially) stall mid-download now and
    # then; a retry almost always succeeds.
    for attempt in range(1, attempts + 1):

        try:
            response = requests.get(url, headers=HEADERS, timeout=timeout)
            response.raise_for_status()
            return response

        except requests.exceptions.RequestException:
            if attempt == attempts:
                raise
            print(f"    retry {attempt} for {url}")
            time.sleep(5 * attempt)


def links(page_url, pattern):
    """Absolute URLs of every href on a page matching pattern."""

    html = get(page_url, timeout=60).text

    found = re.findall(r'href="([^"]+)"', html)

    return [
        urljoin(page_url, href) for href in dict.fromkeys(found)
        if re.search(pattern, href)
    ]


def number(value):

    if value in (None, ""):
        return 0.0

    return float(value)


def row(state, district, fy, amounts, units, source, note=""):

    return {
        "state": state,
        "district": district,
        "fy": fy,
        "cotpa_104": round(amounts.get("104", 0.0), 2),
        "tofei_105": round(amounts.get("105", 0.0), 2),
        "cessation_106": round(amounts.get("106", 0.0), 2),
        "ntcp_activity_total": round(sum(amounts.values()), 2),
        "rop_units": units,
        "mapping_note": note,
        "source_url": source,
    }


def combine(state, fy, units, listing):
    """
    One row per NFHS-5 district from {unit: (district, amounts, url)},
    summing the RoP units that fold into the same district. A row built
    from several RoPs links to the state's listing page.
    """

    rows = []

    for district in sorted({d for d, _, _ in units.values()}):

        names = [u for u, (d, _, _) in units.items() if d == district]

        amounts = {}

        for name in names:
            for code, value in units[name][1].items():
                amounts[code] = amounts.get(code, 0.0) + value

        rows.append(row(
            state, district, fy, amounts, " + ".join(names),
            units[names[0]][2] if len(names) == 1 else listing,
        ))

    return rows


# --- Parsers -----------------------------------------------------------

def parse_haryana_pdf(pdf_bytes):
    """
    {code: lakh} from the "Summary of Approval" pool table, where the
    NCD.4 block reads (layout text):

        NCD.4  National Tobacco Control  104  Implementation of COTPA  6.45
                                         106  Tobacco Cessation        2.20
        NCD.5  ...
    """

    reader = PdfReader(io.BytesIO(pdf_bytes))

    for page in reader.pages[:40]:

        text = page.extract_text(extraction_mode="layout") or ""

        if "NCD.4" not in text:
            continue

        block = text[text.index("NCD.4"):]
        block = block[:block.index("NCD.5")]

        amounts = {}

        for line in block.splitlines():

            match = re.search(
                r"\b(10[4-6])\s+\S.*?\s(\d+\.\d+)\s*$", line
            )

            if match:
                amounts[match.group(1)] = float(match.group(2))

        if not amounts:
            raise ValueError("Haryana RoP: NCD.4 block has no amounts")

        return amounts

    raise ValueError("Haryana RoP: NCD.4 summary not found")


def haryana():

    urls = links(HARYANA_PAGE, r"DistrictROP2024-25/[A-Za-z]+_\d+\.pdf$")

    if len(urls) != 22:
        raise ValueError(f"Haryana: found {len(urls)} of 22 district RoPs")

    rows = []

    for url in urls:

        name = re.search(r"/([A-Za-z]+)_\d+\.pdf$", url).group(1)

        print(f"  Haryana: {name}")

        amounts = parse_haryana_pdf(get(url).content)

        rows.append(row(
            "Haryana", HARYANA_NAMES.get(name, name), "2024-25",
            amounts, name, url,
        ))

    return rows, []


def sikkim():

    url = next(iter(
        links(SIKKIM_PAGE, r"districtwise-rop_2024_25\.xlsx$")
    ), None)

    if url is None:
        raise ValueError("Sikkim: 2024-25 district RoP workbook not found")

    print("  Sikkim: workbook 2024-25")

    book = openpyxl.load_workbook(
        io.BytesIO(get(url).content), data_only=True
    )

    rows = list(book["NCD_Flexipool"].iter_rows(values_only=True))

    # Row 2 opens the "Financial (Lakh)" block; row 3 names its columns.
    start = list(rows[1]).index("Financial (Lakh)")
    names = [str(c).strip() if c else "" for c in rows[2][start:]]

    by_unit = {}
    state_level = 0.0

    for values in rows[3:]:

        code = str(values[2]).strip()

        if code not in CODES:
            continue

        financial = dict(zip(names, values[start:]))

        if number(financial.get("Urban")):
            raise ValueError(f"Sikkim {code}: unexpected urban allocation")

        parts = [number(financial[n]) for n in SIKKIM_NAMES]
        state = number(financial["State"])

        if abs(sum(parts) + state - number(values[6])) > 0.01:
            raise ValueError(f"Sikkim {code}: columns don't sum to total")

        for name in SIKKIM_NAMES:
            by_unit.setdefault(name, {})[code] = number(financial[name])

        state_level += state

    rows_out = []

    for district in dict.fromkeys(SIKKIM_NAMES.values()):

        units = [u for u, d in SIKKIM_NAMES.items() if d == district]

        amounts = {}

        for unit in units:
            for code, value in by_unit[unit].items():
                amounts[code] = amounts.get(code, 0.0) + value

        rows_out.append(row(
            "Sikkim", district, "2024-25", amounts,
            " + ".join(units), url,
        ))

    districts = sum(r["ntcp_activity_total"] for r in rows_out)

    split = [{
        "state": "Sikkim",
        "fy": "2024-25",
        "state_level_lakh": round(state_level, 2),
        "district_level_lakh": round(districts, 2),
        "source_url": url,
    }]

    return rows_out, split


def meghalaya():

    url = next(iter(
        links(MEGHALAYA_PAGE, r"DROP%20NHM%20Meghalaya%202024-26\.xlsx$|"
                              r"DROP NHM Meghalaya 2024-26\.xlsx$")
    ), None)

    if url is None:
        raise ValueError("Meghalaya: 2024-26 district RoP workbook not found")

    url = url.replace(" ", "%20")

    print("  Meghalaya: workbook 2024-26 (sheet 2024-25)")

    book = openpyxl.load_workbook(
        io.BytesIO(get(url).content), data_only=True
    )

    rows = list(book["2024-25"].iter_rows(values_only=True))

    header = next(r for r in rows if "SPMU" in r)
    column = {str(c).strip(): i for i, c in enumerate(header) if c}

    # FMR 104-106 rows inside the NTCP block.
    start = next(
        i for i, r in enumerate(rows)
        if r[1] and "National Tobacco Control" in str(r[1])
    )

    by_code = {}

    for values in rows[start + 1:]:

        code = str(values[0]).strip()

        if code not in CODES:
            if by_code:
                break
            continue

        by_code[code] = values

    if set(by_code) != set(CODES):
        raise ValueError(f"Meghalaya: found NTCP codes {sorted(by_code)}")

    for code, values in by_code.items():

        parts = sum(number(values[column[a]]) for a in MEGHALAYA_NAMES)

        if abs(parts + number(values[column["SPMU"]])
               - number(values[column["Total"]])) > 0.01:
            raise ValueError(f"Meghalaya {code}: columns don't sum")

    rows_out = [
        row(
            "Meghalaya", district, "2024-25",
            {c: number(v[column[abbr]]) for c, v in by_code.items()},
            abbr, url,
        )
        for abbr, district in MEGHALAYA_NAMES.items()
    ]

    split = [{
        "state": "Meghalaya",
        "fy": "2024-25",
        "state_level_lakh": round(sum(
            number(v[column["SPMU"]]) for v in by_code.values()
        ), 2),
        "district_level_lakh": round(
            sum(r["ntcp_activity_total"] for r in rows_out), 2
        ),
        "source_url": url,
    }]

    return rows_out, split


def parse_uttarakhand_pdf(pdf_bytes):
    """
    {code: lakh} from a district RoP. The pool summary near the front
    gives the NTCP main-RoP total, on the same line or the next:

        NCD.4  National Tobacco Control Programme   8.93  0  0  0  0.8  9.73
        NCD.4  National Tobacco Control
                                                   15.92

    and the programme's approval annexure lists each activity:

        NCD 4  National   104   Implementatio   3.38   Rs. 1.60 lakhs ...
                          105   Implementaito   3.40   Rs. 3.40 lakhs ...

    The activity lines must add up to the summary total. In 2026-27 a
    second 104 row (office expenses) belongs to programme management,
    outside the main-RoP total, so only the first row per code is kept.
    """

    reader = PdfReader(io.BytesIO(pdf_bytes))

    total = None
    amounts = {}

    for page in reader.pages:

        text = page.extract_text(extraction_mode="layout") or ""

        if "Tobacco" not in text:
            continue

        lines = text.splitlines()

        for i, line in enumerate(lines):

            if total is None and re.search(
                r"NCD\.\s?4\s+National Tobacco Control", line
            ):
                found = re.search(
                    r"(\d+\.\d+)", " ".join(lines[i:i + 3])
                )
                if found:
                    total = float(found.group(1))

            # Code, activity text, then the first decimal amount.
            match = re.match(
                r"^\s*(?:NCD\.?\s*4\b[^\d]*?)?(10[4-6])\s+"
                r"[A-Za-z].*?\s(\d+\.\d+)\s",
                line,
            )

            if match:
                amounts.setdefault(match.group(1), float(match.group(2)))

    if total is None:
        raise ValueError("Uttarakhand: NCD.4 summary total not found")

    if abs(sum(amounts.values()) - total) > 0.01:
        raise ValueError(
            f"Uttarakhand: NTCP lines {amounts} don't sum to the NCD.4 "
            f"total {total}"
        )

    return amounts


def uttarakhand():

    # The listing is paginated and lays titles out loosely, so every
    # PDF linked near a "RoP_24-25" title is fetched and identified by
    # its cover ("DISTRICT: PAURI GARHWAL / FY 2024-25") instead.
    candidates = []

    for page in range(1, 20):

        url = UTTARAKHAND_PAGE + (f"page/{page}/" if page > 1 else "")

        response = requests.get(url, headers=HEADERS, timeout=60)

        if response.status_code == 404:
            break

        response.raise_for_status()
        html = response.text

        for title in re.finditer(r"RoP_24-25", html):
            window = html[max(0, title.start() - 1500):title.end() + 1500]
            candidates += re.findall(r'href="([^"]+\.pdf)"', window)

    rows = {}

    for url in dict.fromkeys(candidates):

        pdf_bytes = get(url).content

        cover = PdfReader(io.BytesIO(pdf_bytes)).pages[0].extract_text()
        printed = re.search(
            r"DISTRICT:\s*([A-Z][A-Z .]+?)\s*\n\s*FY\s*(\d{4}-\d{2})",
            cover or "",
        )

        if not printed or printed.group(2) != "2024-25":
            continue

        name = printed.group(1).strip().title()
        district = UTTARAKHAND_NAMES.get(name, name)

        if district in rows:
            continue

        print(f"  Uttarakhand: {name}")

        rows[district] = row(
            "Uttarakhand", district, "2024-25",
            parse_uttarakhand_pdf(pdf_bytes), name, url,
        )

    if len(rows) != 13:
        raise ValueError(
            f"Uttarakhand: found {len(rows)} of 13 district RoPs for "
            f"2024-25: {sorted(rows)}"
        )

    return list(rows.values()), []


def parse_odisha_pdf(pdf_bytes):
    """
    {code: lakh} for FY 2024-25 from the "NCD-4_NTCP" approval table:

        Total NTCP                              28.49     33.83
        104  Implementation of COTPA - 2003     24.05     28.97
        105  Implementation of ToEFI Guideline
        106  Tobacco Cessation                   4.44      4.85

    The first amount is 2024-25; a code with no amount was not
    approved. The lines must add up to "Total NTCP".
    """

    reader = PdfReader(io.BytesIO(pdf_bytes))

    total = None
    amounts = {}

    for page in reader.pages:

        text = page.extract_text(extraction_mode="layout") or ""

        if total is None and "Total NTCP" not in text:
            continue

        for line in text.splitlines():

            if total is None:
                found = re.search(r"Total NTCP\s+(\d+\.\d+)", line)
                if found:
                    total = float(found.group(1))
                continue

            match = re.match(
                r"^\s*(10[4-6])\s+[A-Za-z].*?\s(\d+\.\d+)\s", line
            )

            if match:
                amounts.setdefault(match.group(1), float(match.group(2)))

        # 106 is the table's last line; anything further on (targets,
        # annexures) can reuse the same codes for other things.
        if "106" in amounts:
            break

    if total is None:
        raise ValueError("Odisha: Total NTCP row not found")

    # Lines are rounded to two decimals, so their sum can miss the
    # printed total by a paisa or two (Kendrapada: 14.81 vs 14.82).
    if abs(sum(amounts.values()) - total) > 0.02:
        raise ValueError(
            f"Odisha: NTCP lines {amounts} don't sum to Total NTCP {total}"
        )

    return amounts


def odisha():

    # The listing is filled in by a form post (category 193 = 2024-26).
    html = requests.post(
        ODISHA_AJAX,
        data={"categorypips": "193", "action": "myfiltercategorypips"},
        headers=HEADERS, timeout=120,
    ).text

    files = {}

    for url in re.findall(r'href="([^"]+?\.pdf)\s*"', html):

        stem = re.search(r"/([A-Za-z]+)[-_]PIP-2024-26\.pdf$", url)

        if stem and stem.group(1) in ODISHA_NAMES:
            files[stem.group(1)] = url.strip()

    if len(files) != 30:
        raise ValueError(f"Odisha: found {len(files)} of 30 district plans")

    rows = []

    for name, url in sorted(files.items()):

        print(f"  Odisha: {name}")

        rows.append(row(
            "Odisha", ODISHA_NAMES[name], "2024-25",
            parse_odisha_pdf(get(url).content), name, url,
        ))

    return rows, []


def parse_rajasthan_pdf(pdf_bytes):
    """
    {code: lakh} for FY 2024-25 from a district RoP's NCD.4 table, where
    labels wrap and the S.No. column carries the code:

                               Implementation of COTPA -
                        104    2003                       0.60    0.60
                        105    Implementation of ToEFI    3.85    3.85
                        106    Tobacco Cessation          4.50    4.50
                          Total NTCP                      8.95    8.95

    The first amount on a code's line is 2024-25. The lines must add up
    to "Total NTCP".
    """

    reader = PdfReader(io.BytesIO(pdf_bytes))

    amounts = {}

    for page in reader.pages:

        text = page.extract_text(extraction_mode="layout") or ""

        if not re.search(r"COTPA|ToEFI|Tobacco|NTCP", text):
            continue

        for line in text.splitlines():

            match = re.match(r"^.{0,60}?\s(10[4-6])\s+(.*)$", line)

            if match:
                amount = re.search(r"(?<![\d.])(\d+\.\d+)\s", match.group(2) + " ")
                amounts.setdefault(
                    match.group(1), float(amount.group(1)) if amount else 0.0
                )

            total = re.search(r"Total NTCP\s+(\d+\.\d+)", line)

            if total:
                total = float(total.group(1))

                if abs(sum(amounts.values()) - total) > 0.02:
                    raise ValueError(
                        f"Rajasthan: NTCP lines {amounts} don't sum to "
                        f"Total NTCP {total}"
                    )

                return amounts

    raise ValueError("Rajasthan: Total NTCP row not found")


def rajasthan():

    html = get(RAJASTHAN_PAGE, timeout=60).text

    files = {}

    for path in re.findall(r"openPDF\('([^']+\.pdf)'\)", html):

        name = re.sub(r"%20", " ", path.rsplit("/", 1)[1])

        # 2024-26 approvals only (Hanumangarh's is misnamed "2024 23").
        if re.search(r"26-27|letter", name, re.I):
            continue

        stem = re.split(r"\s+(?:NHM|D?RoP)\b", name, flags=re.I)[0].strip()

        if stem not in RAJASTHAN_NAMES:
            raise ValueError(f"Rajasthan: unknown RoP file {name}")

        files[stem] = urljoin(RAJASTHAN_PAGE, path.replace(" ", "%20"))

    if len(files) != len(RAJASTHAN_NAMES):
        raise ValueError(
            f"Rajasthan: found {len(files)} of {len(RAJASTHAN_NAMES)} "
            f"district RoPs; missing {sorted(set(RAJASTHAN_NAMES) - set(files))}"
        )

    units = {}

    for stem, url in sorted(files.items()):

        if stem in RAJASTHAN_UNREADABLE:
            print(f"  Rajasthan: {stem} skipped (no amounts in file)")
            continue

        print(f"  Rajasthan: {stem}")

        units[stem] = (
            RAJASTHAN_NAMES[stem], parse_rajasthan_pdf(get(url).content), url
        )

    rows = combine("Rajasthan", "2024-25", units, RAJASTHAN_PAGE)

    for r in rows:
        r["mapping_note"] = rajasthan_note(r["district"])

    return rows, []


def rajasthan_note(district):
    """Which split 2023 districts this NFHS-5 district gains or loses."""

    notes = []

    for unit, shares in RAJASTHAN_SPLITS.items():

        if district not in shares:
            continue

        home = RAJASTHAN_NAMES[unit]
        others = ", ".join(
            f"{share:.0%} {d}" for d, share in shares.items() if d != home
        )

        if district == home:
            notes.append(
                f"includes all of {unit} ({shares[home]:.0%} from "
                f"{home}; {others})"
            )
        else:
            notes.append(
                f"excludes {unit} ({shares[district]:.0%} from "
                f"{district}; counted under {home})"
            )

    return "; ".join(notes)


def assam():
    """
    The NTCP sheet is one wide table split over 14 pages. Reading left
    to right across pages, every row has: fresh approval (2 columns),
    5 text columns, state HQ (6), then for each of 34 districts a basis
    column and quantity / unit cost / amount for 2024-25 and 2025-26
    (7), then State Total and "Total fund allocated to the districts"
    (2 each). The "Programme Total" row fills every column, so it fixes
    the column positions on each page; activity lines are read against
    them and must add up to both the district and the sheet totals.
    """

    print("  Assam: NTCP district budget sheet")

    pdf = pdfplumber.open(io.BytesIO(get(ASSAM_SHEET).content))

    numeric = re.compile(r"^\d+(\.\d+)?$")

    columns = {}         # page -> [(x centre, sheet column)]
    totals = []
    names = {}
    lines = {}           # code -> {sheet column: value}

    for number_, page in enumerate(pdf.pages):

        words = page.extract_words(x_tolerance=1.5)

        label = next(w for w in words if w["text"] == "Total:")

        for w in sorted(
            (w for w in words if abs(w["top"] - label["top"]) < 3
             and w["x0"] > label["x1"] and numeric.match(w["text"])),
            key=lambda w: w["x0"],
        ):
            columns.setdefault(number_, []).append(
                ((w["x0"] + w["x1"]) / 2, len(totals))
            )
            totals.append(float(w["text"]))

        heading = {}
        for w in words:
            if w["top"] < 70:
                heading.setdefault(round(w["top"]), []).append(w)

        for ws in heading.values():
            text = " ".join(
                w["text"] for w in sorted(ws, key=lambda w: w["x0"])
            )
            names.update({
                int(n): name.strip() for n, name in re.findall(
                    r"(\d+)\s*-\s*Fund allocated to (.+?)"
                    r"(?=\s+\d+\s*-\s*Fund|\s+State Total|\s+Total fund|$)",
                    text,
                )
            })

        # Values sit on the same baseline as their row's code, to the
        # right of the label columns.
        for code in (w for w in words if w["x0"] < 100):

            found = re.match(r"^(10[46]\.\d+)", code["text"])

            if not found:
                continue

            for w in words:
                if (abs(w["top"] - code["top"]) < 3 and w["x0"] > 330
                        and numeric.match(w["text"])):
                    x = (w["x0"] + w["x1"]) / 2
                    gap, column = min(
                        (abs(x - cx), c) for cx, c in columns[number_]
                    )
                    if gap < 16:
                        lines.setdefault(found.group(1), {})[column] = (
                            float(w["text"])
                        )

    count = len(names)
    lead = 13

    if sorted(names) != list(range(1, count + 1)) or (
            len(totals) != lead + 7 * count + 4):
        raise ValueError(
            f"Assam: {count} district headings and {len(totals)} total "
            "columns don't fit the expected layout"
        )

    def amount_2024(values, i):
        return values.get(lead + 7 * i + 3, 0.0)

    district_fund = len(totals) - 2      # "Total fund ... 2024-25"

    for code, values in lines.items():
        spread = sum(amount_2024(values, i) for i in range(count))
        if abs(spread - values.get(district_fund, 0.0)) > 0.011:
            raise ValueError(
                f"Assam {code}: districts sum to {spread:.2f}, sheet says "
                f"{values.get(district_fund, 0.0):.2f}"
            )

    units = {}

    for i in range(count):

        amounts = {"104": 0.0, "105": 0.0, "106": 0.0}

        for code, values in lines.items():
            key = (
                "104" if code.startswith("104")
                else "105" if code in ASSAM_TOFEI else "106"
            )
            amounts[key] += amount_2024(values, i)

        if abs(sum(amounts.values()) - amount_2024(
                dict(enumerate(totals)), i)) > 0.011:
            raise ValueError(f"Assam: {names[i + 1]} lines miss its total")

        name = names[i + 1]
        units[name] = (ASSAM_NAMES.get(name, name), amounts, ASSAM_SHEET)

    rows = combine("Assam", "2024-25", units, ASSAM_SHEET)

    state_level = totals[lead - 4]       # state HQ, 2024-25 amount
    districts = sum(r["ntcp_activity_total"] for r in rows)

    if abs(state_level + districts - totals[0]) > 0.02:
        raise ValueError(
            "Assam: state HQ plus districts doesn't match the fresh "
            "2024-25 approval"
        )

    split = [{
        "state": "Assam",
        "fy": "2024-25",
        "state_level_lakh": round(state_level, 2),
        "district_level_lakh": round(districts, 2),
        "source_url": ASSAM_SHEET,
    }]

    return rows, split


STATES = [
    haryana, sikkim, meghalaya, uttarakhand, odisha, rajasthan, assam,
]


def write(path, rows):

    with open(path, "w", encoding="utf-8", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def run():

    districts, splits = [], []

    for state in STATES:
        rows, split = state()
        districts += rows
        splits += split

    OUT_DIR.mkdir(parents=True, exist_ok=True)

    write(DISTRICT_PATH, districts)
    write(SPLIT_PATH, splits)

    print(
        f"✓ NTCP district RoPs: {len(districts)} districts -> "
        f"{DISTRICT_PATH.name}, {SPLIT_PATH.name}"
    )


if __name__ == "__main__":
    run()
